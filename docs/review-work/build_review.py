from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = Path(__file__).resolve().parents[1] / 'AIVES_Document_Review.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
sec.footer_distance = Cm(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
    for borders in s._element.findall('.//' + qn('w:pBdr')):
        borders.getparent().remove(borders)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.13
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(6)
for s in ['List Bullet','List Number']:
    doc.styles[s].font.size=Pt(10.5)
    doc.styles[s].paragraph_format.space_after=Pt(6)
    doc.styles[s].paragraph_format.line_spacing=1.1
footer = sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('AIVES  |  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('666666')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Tài liệu review dự án AIVES'
doc.core_properties.subject='Nghiệp vụ và mức độ triển khai hệ thống thi vấn đáp có AI'
doc.core_properties.author='Nhóm dự án AIVES'

def p(text, style=None): return doc.add_paragraph(text,style)
def h(text): return doc.add_heading(text,2)
def bullet(text): return p(text,'List Bullet')
current_num = None
def step(text):
    global current_num
    if current_num is None:
        numbering = doc.part.numbering_part.element
        original = numbering.find("w:num[@w:numId='5']", numbering.nsmap)
        abstract = original.find(qn('w:abstractNumId')).get(qn('w:val'))
        current_num = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
        num = OxmlElement('w:num'); num.set(qn('w:numId'),str(current_num))
        aid = OxmlElement('w:abstractNumId'); aid.set(qn('w:val'),abstract);num.append(aid)
        override = OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0')
        start = OxmlElement('w:startOverride');start.set(qn('w:val'),'1');override.append(start);num.append(override)
        numbering.append(num)
    para=p(text,'List Number')
    np=para._p.get_or_add_pPr().get_or_add_numPr();np.get_or_add_ilvl().val=0;np.get_or_add_numId().val=current_num
    return para
def page(title):
    global current_num
    current_num = None
    doc.add_page_break()
    doc.add_heading(title,1)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for cell,w,text in zip(t.rows[0].cells,widths,headers): cell.width=Cm(w);cell.text=text
    for row in rows:
        cells=t.add_row().cells
        for c,w,text in zip(cells,widths,row): c.width=Cm(w);c.text=text
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+name);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0: pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF3' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'));props.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'90');m.set(qn('w:type'),'dxa');margins.append(m)
            props.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.07
                for run in para.runs: run.font.size=Pt(9.5);run.bold=(i==0)
    p('')
    return t

doc.add_heading('Tài liệu review dự án AIVES',0)
p('Hệ thống thi vấn đáp thông minh có sử dụng AI', 'Subtitle')
p('AI powered Viva Exam System  |  SWD392  |  Ngày 18 tháng 10 năm 2026')
h('Kết luận review')
p('AIVES đã triển khai các luồng chính để giảng viên chuẩn bị câu hỏi và rubric, tổ chức kỳ thi, và để sinh viên tham gia phỏng vấn bằng giọng nói hoặc văn bản. Điểm trọng tâm của dự án là AI hỏi làm rõ dựa trên nội dung câu trả lời, trong giới hạn do giảng viên thiết lập. Chuỗi đánh giá sau thi còn thiếu chấm điểm AI, màn hình giảng viên duyệt điểm và báo cáo kết quả.')
h('Vấn đề cần giải quyết')
p('Thi vấn đáp cần nhiều thời gian của giảng viên; cách hỏi và thang chấm có thể khác nhau giữa các ca thi. Khi số lượng sinh viên tăng, việc phân câu hỏi, theo dõi câu trả lời và lưu bằng chứng để giải thích kết quả trở nên khó khăn. AIVES hỗ trợ chuẩn bị đề, điều phối hội thoại và lưu lại dữ liệu phục vụ đánh giá.')
h('Core của dự án')
p('Lõi của AIVES là vòng phỏng vấn thích ứng: đặt câu hỏi, nhận câu trả lời, chuyển giọng nói thành văn bản, dùng AI xác định điểm cần làm rõ, rồi đặt câu đào sâu hoặc chuyển câu chính tiếp theo. Giảng viên kiểm soát nguồn câu hỏi, rubric, lịch thi và giới hạn số câu; AI vận hành hội thoại trong phạm vi đó.')
bullet('Ngân hàng câu hỏi chỉ tiếp nhận câu đã được giảng viên duyệt; AI không có quyền tự phê duyệt.')
bullet('Câu hỏi có mức Bloom và rubric để giữ căn cứ đánh giá nhất quán giữa ra đề, phỏng vấn và chấm điểm về sau.')
bullet('Điểm cuối cùng thuộc quyết định của giảng viên trong nghiệp vụ mục tiêu; luồng chốt điểm hiện chưa được triển khai.')
h('Phạm vi tài liệu')
p('Tài liệu tập trung review nhóm chức năng 1, 2 và 3, kiến trúc đang dùng, kịch bản demo và phần cần hoàn thiện. Trạng thái “đã triển khai” mô tả mã nguồn hiện có; nghiệm thu vận hành cần thêm một lượt thi tích hợp trên môi trường triển khai.')

page('1 Vai trò và quản lý ngân hàng câu hỏi')
table(['Vai trò','Trách nhiệm'],[
('Quản trị viên','Quản lý tài khoản, môn học và phân công giảng viên.'),
('Giảng viên','Quản lý ghi danh, tài liệu, rubric, câu hỏi và kỳ thi của môn được phân công.'),
('Sinh viên','Xem lịch của mình, vào thi đúng lịch, trả lời và xem lại hội thoại đã nộp.')],[3.4,13.6])
h('Luồng chuẩn bị câu hỏi')
step('Giảng viên tạo rubric nháp, thêm tiêu chí và điểm tối đa cho từng tiêu chí. Hệ thống chỉ cho công bố khi tổng điểm tiêu chí bằng điểm tối đa của rubric.')
step('Giảng viên chọn một trong ba cách: import CSV hoặc JSON, nhờ AI sinh từ tài liệu, hoặc tạo thủ công. Mỗi câu được gắn Bloom, độ khó và rubric phù hợp.')
step('Câu mới đi vào PENDING_REVIEW. Giảng viên xem nội dung, đáp án tham khảo, tiêu chí chấm và đoạn tài liệu nguồn nếu câu do AI sinh.')
step('Giảng viên quyết định chỉnh sửa, giữ chờ, từ chối, loại bỏ hoặc duyệt. Chỉ APPROVED mới thuộc ngân hàng chính thức và được thêm vào đề.')
h('Quy tắc bảo vệ chất lượng')
bullet('Bloom hỗ trợ bốn mức: Nhớ, Hiểu, Vận dụng và Phân tích. Độ khó được lưu riêng với các mức EASY, MEDIUM, HARD.')
bullet('Sửa câu đã duyệt đưa câu về chờ duyệt và gỡ khỏi đề nháp. Câu đang dùng trong kỳ thi ACTIVE hoặc SCHEDULED bị khóa sửa và loại bỏ.')
bullet('Loại bỏ chuyển câu sang ARCHIVED, giữ lịch sử thay vì xóa vật lý. Hệ thống ghi người duyệt, thời điểm và quyết định review.')
h('Import và sinh câu hỏi bằng AI')
p('Import hỗ trợ xem trước và lỗi từng dòng; tối đa 2 MB và 200 câu mỗi lần. Commit kiểm tra lại dữ liệu; có dòng lỗi thì không lưu một phần. Hệ thống kiểm tra nội dung trùng và rubric đã công bố thuộc đúng môn.')
p('Luồng AI nhận tối đa 5 tài liệu PDF hoặc DOCX đã xử lý, tổng nguồn không quá 60.000 ký tự, sinh 1–20 câu mỗi lần. Kết quả có đáp án và liên kết đoạn nguồn. Giao diện hiển thị trạng thái xử lý, lỗi tại form, kết quả và lịch sử để giảng viên tiếp tục duyệt.')
p('Đánh giá: luồng giảng viên kiểm soát ngân hàng câu hỏi đã có. Cần nghiệm thu thêm chất lượng câu hỏi thực tế, độ đúng của Bloom và mức độ bám sát rubric; cấu trúc JSON hợp lệ chưa chứng minh nội dung học thuật chính xác.')

page('2 Tổ chức kỳ thi và chọn bộ câu hỏi')
h('Luồng nghiệp vụ của giảng viên')
step('Tạo kỳ thi DRAFT cho môn được phân công; đặt số câu chính, tổng số câu đào sâu tối đa mỗi sinh viên, chế độ chọn câu và thời gian trả lời.')
step('Thêm câu APPROVED vào ngân hàng đề của kỳ thi. Hệ thống kiểm tra đủ số câu hợp lệ trước khi kích hoạt.')
step('Chọn sinh viên đã ghi danh hoặc nhập email tài khoản STUDENT đang hoạt động. Có tùy chọn ghi danh vào môn và phân lịch trong cùng giao dịch.')
step('Kiểm tra danh sách thí sinh và khung giờ; được sửa hoặc bỏ lịch trước khi có lượt thi. Sinh viên chỉ bắt đầu trong lịch của mình và khi kỳ thi ACTIVE.')
step('Kết thúc kỳ thi khi không còn lượt IN_PROGRESS. Kỳ thi CLOSED giữ lịch sử và không nhận lượt bắt đầu mới.')
h('Cơ chế phân câu hỏi')
table(['Chế độ','Hành vi hiện tại'],[
('RANDOM','Trộn ngẫu nhiên các câu hợp lệ; ưu tiên tránh những câu đã dùng trong lượt bắt đầu ngay trước của cùng kỳ thi.'),
('ADAPTIVE','Ngoài tránh bộ câu trước, ưu tiên câu ít xuất hiện trong 20 lượt gần nhất và đa dạng mức Bloom, độ khó.')],[3.4,13.6])
p('ADAPTIVE ở bước cấp bộ câu dựa trên lịch sử sử dụng đề. Cơ chế hiện tại chưa ước lượng năng lực sinh viên từ điểm số. Việc hỏi đào sâu trong phòng thi là một bước thích ứng riêng, dựa trên câu trả lời.')
h('Giới hạn và tình huống biên')
bullet('Ví dụ 3 câu chính và tối đa 4 câu đào sâu cho một sinh viên thì toàn bộ lượt thi có tối đa 7 câu. Ngân sách đào sâu không đặt lại khi chuyển câu chính hoặc tải lại trang.')
bullet('Nếu có đủ ít nhất hai lần số câu chính, hệ thống có thể tránh hoàn toàn bộ câu ngay trước. Nếu ít hơn, phần trùng tối thiểu là max của 0 và hai lần số câu chính trừ số câu hợp lệ.')
bullet('Hệ thống khóa theo kỳ thi khi cấp câu để xử lý các sinh viên bắt đầu đồng thời. Mỗi lượt lưu snapshot nội dung câu hỏi, đáp án và rubric để bảo toàn căn cứ về sau.')
p('Đánh giá: giảng viên đã có luồng tạo đề, quản lý thí sinh và lịch thi. Khả năng tránh trùng phụ thuộc quy mô ngân hàng đề; cần chuẩn bị đủ câu cho buổi demo và đánh giá tải khi nhiều sinh viên cùng bắt đầu.')

page('3 Phòng phỏng vấn AI')
h('Vòng phỏng vấn đang triển khai')
step('Sinh viên bắt đầu lượt thi hợp lệ. Backend cấp một câu hiện tại và giữ kín đáp án tham khảo cùng rubric riêng của đề.')
step('Hệ thống dùng text to speech để đọc câu hỏi. Sinh viên có thể nghe lại và chọn ghi âm hoặc bắt đầu trả lời bằng văn bản.')
step('Khi bắt đầu trả lời, backend ghi deadline. Ghi âm được chuyển thành văn bản trực tiếp qua WebSocket; khi nộp, backend xử lý file âm thanh để lưu transcript chính thức.')
step('AI nhận câu hỏi, đáp án và rubric snapshot cùng hội thoại; quyết định hỏi làm rõ hoặc chuyển tiếp. Câu đào sâu phải nằm trong tổng ngân sách của sinh viên và thời gian ca thi.')
step('Khi hoàn tất hoặc hết giờ, hệ thống lưu trạng thái lượt thi và hội thoại đã nộp. Sinh viên xem lại hội thoại và nghe bản ghi đã lưu của chính mình.')
h('Ví dụ minh họa hành vi mong muốn')
p('Câu chính: “Khi nào nên dùng kiến trúc microservices?” Sinh viên trả lời: “Khi hệ thống lớn vì dễ mở rộng.” AI có thể hỏi: “Nếu nhóm chỉ có ba người và cần giao hàng nhanh, em cân nhắc chi phí vận hành như thế nào?” Mục đích là làm rõ lập luận của sinh viên; nội dung hỏi tiếp cụ thể phụ thuộc kết quả AI.')
h('Kiểm soát và khả năng phục hồi')
bullet('Deadline do backend quản lý, không đặt lại khi reload và không vượt thời điểm kết thúc lịch thí sinh. Worker kiểm tra lượt đang thi định kỳ.')
bullet('Khóa lượt thi và promptId giúp tránh nộp trùng hoặc nộp câu trả lời cũ sang câu hỏi mới khi có nhiều yêu cầu đồng thời.')
bullet('Nếu sinh câu đào sâu thất bại, câu trả lời đã nộp vẫn được giữ. Thử tiếp xử lý AI không gửi trùng câu trả lời đã lưu.')
bullet('API key chỉ ở backend. WebSocket dùng vé ngắn hạn một lần; quyền truy cập ràng buộc với sinh viên và câu đang trả lời.')
h('Giới hạn vận hành hiện tại')
p('Micro cần localhost hoặc HTTPS và quyền trình duyệt. Âm thanh đang ghi nhưng chưa nộp có thể mất khi đóng hoặc tải lại trang. Transcript trực tiếp là bản tạm; transcript từ file nộp là bản chính thức. Chưa có ghi hình liên tục, giám sát chống gian lận hoặc chấm điểm AI.')
p('Đánh giá: vòng phỏng vấn thích ứng là phần lõi đã có mã nguồn và kiểm thử. Cần chạy thử đầu cuối với micro thật, mạng thực tế và tài liệu môn học trước khi nghiệm thu buổi thi.')

page('4 Kiến trúc và dữ liệu')
h('Các thành phần chính')
table(['Thành phần','Công nghệ và vai trò'],[
('Giao diện','React, Vite và Tailwind CSS; màn hình riêng theo vai trò, gọi REST API và WebSocket backend.'),
('Backend','Java 21, Spring Boot 3.4.1, Spring Security JWT, Spring Data JPA; xử lý nghiệp vụ và phân quyền theo môn.'),
('Cơ sở dữ liệu','PostgreSQL trên Neon; lưu tài khoản, quyền môn học, rubric, câu hỏi, lịch thi và hội thoại.'),
('Lưu trữ tệp','Ổ đĩa riêng của backend cho tài liệu gốc và âm thanh. Cần vùng lưu trữ bền vững khi triển khai.'),
('Dịch vụ AI','OpenAI cho sinh câu hỏi, quyết định hỏi đào sâu, nhận dạng giọng nói và tạo giọng đọc.')],[3.6,13.4])
h('Đường đi của dữ liệu')
p('Trình duyệt gửi yêu cầu đã xác thực đến backend. Backend kiểm tra vai trò, phân công môn hoặc quyền sở hữu lượt thi trước khi đọc và ghi Neon. Các yêu cầu AI được gọi từ backend; trình duyệt không kết nối database hoặc giữ API key của nhà cung cấp.')
p('Tài liệu đi qua bước trích xuất và chia đoạn trước khi sinh câu hỏi. Mỗi câu AI liên kết với đoạn nguồn. Cơ chế hiện tại đưa văn bản của các tài liệu được chọn vào ngữ cảnh trong giới hạn ký tự; chưa có tìm kiếm vector hoặc truy hồi ngữ nghĩa để xử lý kho tài liệu lớn.')
h('Nhóm dữ liệu phục vụ truy vết')
bullet('Nguồn học liệu: course_materials và material_chunks lưu metadata, trạng thái xử lý và văn bản nguồn; tệp gốc nằm trên ổ đĩa.')
bullet('Ngân hàng: rubrics, rubric_criteria, questions, question_reviews, question_generation_jobs và question_import_batches lưu chuẩn chấm, câu hỏi và quá trình tạo, duyệt.')
bullet('Kỳ thi: exams, exam_question_pool, exam_candidates, exam_attempts và attempt_questions lưu cấu hình, phân lịch và snapshot đề của từng lượt.')
bullet('Hội thoại: interview_turns và ai_interview_requests lưu lượt hỏi đáp, liên kết âm thanh và trạng thái xử lý AI.')
h('Điểm cần củng cố khi triển khai')
p('Cần quản lý thay đổi schema bằng migration có phiên bản, sao lưu đồng bộ database và tệp, cấu hình HTTPS cùng proxy WebSocket. Đồng thời cần chính sách lưu giữ và quyền truy cập dữ liệu thi. Các cột văn bản dài phải giữ kiểu TEXT để tránh lỗi giới hạn 255 ký tự.')

page('5 Mức độ hoàn thành và phần còn thiếu')
p('Bảng dưới đây phân biệt khả năng đã có trong mã nguồn với chuỗi nghiệp vụ còn cần phát triển. Không quy đổi thành phần trăm khi chưa có bộ tiêu chí nghiệm thu và kết quả chạy tích hợp đầy đủ.')
table(['Nhóm','Trạng thái','Phần cần hoàn thiện'],[
('1 Câu hỏi và rubric','Đã triển khai luồng chính','Nghiệm thu nội dung AI, chất lượng nguồn và Bloom; bổ sung truy hồi ngữ nghĩa nếu mở rộng kho tài liệu.'),
('2 Kỳ thi và lịch','Đã triển khai luồng chính','Chạy tích hợp nhiều thí sinh; kiểm tra tính công bằng và tải khi bắt đầu đồng thời.'),
('3 Phỏng vấn AI','Đã triển khai lõi hội thoại','Thử micro thật, mạng chậm, mất kết nối và chi phí AI trong buổi thi đầy đủ.'),
('4 Chấm điểm hỗ trợ AI','Chưa hoàn thiện','Chấm từng tiêu chí, giải thích kèm bằng chứng, giảng viên sửa và chốt điểm.'),
('5 Minh bạch và giám sát','Đã có một phần dữ liệu','Có transcript và audio đã nộp; thiếu dashboard giảng viên, ghi hình và quản lý bằng chứng toàn buổi.'),
('6 Phản hồi và báo cáo','Chưa hoàn thiện','Công bố kết quả, báo cáo sinh viên, thống kê lớp và xuất bảng điểm.')],[3.6,4.1,9.3])
h('Khoảng trống quan trọng nhất')
p('Phỏng vấn hiện tạo được bằng chứng đầu vào cho đánh giá, nhưng chưa tạo ra kết quả học tập đã được giảng viên phê duyệt. Snapshot rubric và transcript là nền tảng cho bước chấm điểm; sự tồn tại của hai loại dữ liệu này chưa đồng nghĩa hệ thống đã chấm điểm được.')
h('Các tiêu chí nghiệm thu còn cần chứng minh')
bullet('Chất lượng học thuật: câu AI bám nguồn, Bloom phù hợp, câu đào sâu có ích và không dẫn dắt đáp án.')
bullet('Tính ổn định: không mất câu trả lời đã nộp, không vượt giờ hoặc ngân sách đào sâu, không cấp nhầm dữ liệu giữa sinh viên.')
bullet('Khả năng vận hành: tệp tồn tại sau restart hoặc triển khai lại; chi phí và độ trễ AI đáp ứng quy mô ca thi.')
bullet('Tính công bằng: tránh trùng câu không tự đảm bảo độ khó tương đương. Cần kiểm tra độ phủ rubric và mức khó giữa các bộ đề.')

page('6 Kịch bản demo và bằng chứng kiểm thử')
h('Kịch bản demo xuyên suốt')
step('Admin chuẩn bị tài khoản và phân công môn. Giảng viên đăng nhập, kiểm tra môn và ghi danh sinh viên thử nghiệm.')
step('Tạo rubric 10 điểm, thêm tiêu chí đủ tổng 10 và công bố. Thử công bố rubric thiếu tổng điểm để chứng minh kiểm soát nghiệp vụ.')
step('Upload PDF hoặc DOCX có văn bản, xử lý đến READY. Chọn tài liệu, rubric đã công bố và sinh 2–3 câu; kiểm tra kết quả và nguồn.')
step('Import một file mẫu, xem trước lỗi nếu có. Cho thấy câu mới vẫn ở chờ duyệt; sửa một câu, từ chối một câu và duyệt các câu phù hợp.')
step('Tạo kỳ thi 2 câu chính và tổng tối đa 2 câu đào sâu, thêm đủ câu đã duyệt, phân lịch hai sinh viên rồi kích hoạt.')
step('Sinh viên thứ nhất vào thi, nghe câu hỏi, trả lời bằng micro hoặc text. Quan sát transcript, câu đào sâu nếu AI yêu cầu và bộ đếm giới hạn.')
step('Tải lại trang để kiểm tra trạng thái và deadline được giữ; hoàn tất lượt thi rồi mở hội thoại đã lưu. Sinh viên tiếp theo bắt đầu để đối chiếu bộ câu.')
h('Bằng chứng kiểm tra hiện có')
table(['Phạm vi','Bằng chứng và giới hạn'],[
('Backend','Báo cáo Surefire hiện có tổng 55 test, không failure hoặc error. Các suite chạy ở những lần khác nhau; lần gần nhất tập trung TeachingFeaturesTests gồm 7 test.'),
('Frontend','Lần xác minh gần nhất ghi nhận 14 test API, 23 test Playwright và build Vite thành công. Playwright dùng API và micro giả lập.'),
('Dịch vụ thật','Đã thử riêng các API phỏng vấn OpenAI bằng âm thanh tổng hợp và câu mẫu; chưa có bằng chứng một buổi thi hoàn chỉnh bằng dữ liệu sinh viên thật.')],[3.6,13.4])
h('Phạm vi cần chạy trước buổi review')
p('Khởi động lại backend để nạp API mới và tải lại frontend. Kiểm tra health, JWT, phân công môn, rubric PUBLISHED, tài liệu READY, quyền micro và kết nối WebSocket. Dùng tài khoản sinh viên thử nghiệm trong đúng khung giờ. Kiểm thử tự động chưa thay thế việc nghiệm thu trực tiếp chất lượng âm thanh và hội thoại AI.')
p('Nếu sinh câu hỏi chưa chạy, kiểm tra lựa chọn rubric tại form và thông báo ngay cạnh nút sinh; nếu provider lỗi, xem trạng thái FAILED trong lịch sử. Không coi thông báo đã cấu hình API key là bằng chứng yêu cầu AI đã thành công.')

page('7 Ưu tiên phát triển và nội dung trao đổi review')
h('Thứ tự hoàn thiện đề xuất')
table(['Ưu tiên','Hạng mục','Điều kiện hoàn thành'],[
('1','Ổn định nhóm 1 đến 3','Chạy một lượt thi đầu cuối trên môi trường thật; xác minh lưu tệp, giới hạn thời gian, quyền và phục hồi lỗi.'),
('2','Chấm điểm AI theo rubric','Điểm từng tiêu chí có giới hạn, nhận xét và dẫn chứng từ transcript; lưu phiên bản model, rubric và kết quả xử lý.'),
('3','Giảng viên duyệt điểm','Giảng viên xem hội thoại, đối chiếu đề, chỉnh điểm, ghi lý do và chốt; lưu lịch sử quyết định.'),
('4','Phản hồi và báo cáo','Sinh viên chỉ xem kết quả đã công bố; giảng viên xem thống kê lớp và xuất bảng điểm.')],[1.7,4.5,10.8])
h('Các câu hỏi thường gặp khi review')
p('AI có tự quyết định câu nào vào ngân hàng không? Không. Mọi câu mới chờ giảng viên kiểm tra và duyệt; chỉ câu APPROVED được dùng trong đề.')
p('Tính thích ứng nằm ở đâu? Bước cấp đề cân bằng lịch sử sử dụng câu và độ phủ Bloom, độ khó. Trong hội thoại, AI dựa vào câu trả lời để quyết định hỏi làm rõ. Hai cơ chế này phục vụ hai mục đích khác nhau.')
p('Có đảm bảo hai sinh viên không trùng câu không? Hệ thống ưu tiên tránh bộ câu ngay trước; chỉ tránh hoàn toàn khi nguồn câu đủ. Giới hạn này được thể hiện bằng cảnh báo tại màn hình kỳ thi.')
p('Hiện tại ai chấm và chốt điểm? Đây là bước chưa hoàn thiện trong sản phẩm. Nghiệp vụ mục tiêu là AI gợi ý theo rubric và giảng viên quyết định điểm cuối cùng; không nên demo hoặc giới thiệu như chức năng đã có.')
p('Dự án khác một hệ thống trắc nghiệm ở điểm nào? Câu trả lời mở và hội thoại nhiều lượt cho phép AI yêu cầu sinh viên giải thích lập luận. Transcript và âm thanh đã nộp được giữ để phục vụ xem xét sau thi.')
h('Mã nguồn và tài liệu đối chiếu')
p('Frontend: src/components/api/Questions.jsx, QuestionImport.jsx, Materials.jsx, Exams.jsx; src/components/Interview.jsx và StudentWorkspace.jsx. Backend: QuestionGenerationService, QuestionSetSelector, InterviewStateStore, TeacherExamManagementController và các test nghiệp vụ.')
p('Tài liệu dự án: TEACHER_WORKFLOW.md, TEACHING_FEATURES.md và INTERVIEW_FEATURE.md tại thư mục backend SWD392_SE1907_AIVES. Báo cáo kiểm thử: target/surefire-reports của backend và test-results của frontend.')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
